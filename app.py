# REQUIREMENTS.TXT

```txt id="n6p2zi"
streamlit
pandas
gspread
google-auth
openpyxl
python-docx
reportlab
Pillow
```

# APP.PY — FULL PERFECT PROFESSIONAL VERSION 2026

```python id="44cf9i"
# =========================================================
# BUKU TAMU DIGITAL BAPPEDA KOTA PARIAMAN
# FULL PROFESSIONAL VERSION 2026
# FINAL STABLE VERSION
# =========================================================

# =========================================================
# IMPORT
# =========================================================

import streamlit as st
import pandas as pd
import gspread
import base64
import re

from io import BytesIO
from datetime import datetime
from google.oauth2.service_account import Credentials

# =========================================================
# OPTIONAL IMPORT DOCX
# =========================================================

try:
    from docx import Document
    DOCX_AVAILABLE = True
except:
    DOCX_AVAILABLE = False

# =========================================================
# OPTIONAL IMPORT PDF
# =========================================================

try:

    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        PageBreak
    )

    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import letter

    PDF_AVAILABLE = True

except:
    PDF_AVAILABLE = False

# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="Buku Tamu Digital BAPPEDA",
    page_icon="📘",
    layout="wide",
    initial_sidebar_state="expanded"
)

# =========================================================
# CUSTOM CSS
# =========================================================

st.markdown("""
<style>

html, body, [class*="css"]  {
    font-family: 'Segoe UI', sans-serif;
}

.main {
    background-color: #F4F7FE;
}

section[data-testid="stSidebar"] {
    background: linear-gradient(
        180deg,
        #081F4D 0%,
        #0D2F75 100%
    );
}

section[data-testid="stSidebar"] * {
    color: white !important;
}

.sidebar-title {
    font-size: 30px;
    font-weight: 700;
    margin-bottom: 5px;
}

.sidebar-subtitle {
    font-size: 15px;
    color: #D1D5DB !important;
    margin-bottom: 25px;
}

.stButton > button {
    width: 100%;
    border: none;
    border-radius: 10px;
    padding: 0.8rem;
    font-weight: bold;
    color: white !important;
    background: linear-gradient(
        90deg,
        #2563EB,
        #1D4ED8
    );
}

.stDownloadButton > button {
    width: 100%;
    border: none;
    border-radius: 10px;
    padding: 0.8rem;
    font-weight: bold;
    color: white !important;
    background: linear-gradient(
        90deg,
        #10B981,
        #059669
    );
}

.stMetric {
    background: white;
    padding: 15px;
    border-radius: 12px;
    box-shadow: 0px 3px 10px rgba(0,0,0,0.05);
}

[data-testid="stDataFrame"] {
    border-radius: 12px;
}

@media (max-width: 768px) {

    .block-container {
        padding: 1rem;
    }

    .stButton button {
        font-size: 14px;
    }

}

</style>
""", unsafe_allow_html=True)

# =========================================================
# GOOGLE SHEETS CONNECTION
# =========================================================

@st.cache_resource
def connect_gsheet():

    scope = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]

    creds_dict = st.secrets["gcp_service_account"]

    creds = Credentials.from_service_account_info(
        creds_dict,
        scopes=scope
    )

    client = gspread.authorize(creds)

    spreadsheet = client.open_by_key(
        "1lBGe8ZTLBICZz5dbDgPqwNiv4FO-CEFmcSnczYNUxz8"
    )

    sheet = spreadsheet.sheet1

    return sheet

try:

    sheet = connect_gsheet()

except Exception as e:

    st.error(
        f"Gagal koneksi Google Sheets: {e}"
    )

    st.stop()

# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.markdown("""
<div class="sidebar-title">
📘 Buku Tamu Digital
</div>

<div class="sidebar-subtitle">
BAPPEDA Kota Pariaman
</div>
""", unsafe_allow_html=True)

menu = st.sidebar.radio(
    "📌 Menu",
    [
        "Dashboard",
        "Input Buku Tamu",
        "Daftar Buku Tamu",
        "Ringkasan Statistik"
    ]
)

# =========================================================
# LOAD DATA
# =========================================================

@st.cache_data(ttl=30)
def load_data():

    try:

        data = sheet.get_all_values()

        columns = [
            "tanggal",
            "nama",
            "opd",
            "nomor_hp",
            "bidang",
            "foto",
            "spt"
        ]

        if len(data) <= 1:

            return pd.DataFrame(
                columns=columns
            )

        headers = [
            h.lower().strip()
            for h in data[0]
        ]

        df = pd.DataFrame(
            data[1:],
            columns=headers
        )

        for col in columns:

            if col not in df.columns:
                df[col] = ""

        return df

    except Exception as e:

        st.error(
            f"Gagal memuat data: {e}"
        )

        return pd.DataFrame()

# =========================================================
# IMAGE DECODER
# =========================================================

def show_base64_image(base64_string, caption):

    try:

        if (
            pd.isna(base64_string)
            or str(base64_string).strip() == ""
        ):

            st.warning(
                "Gambar tidak tersedia."
            )

            return

        image_bytes = base64.b64decode(
            str(base64_string)
        )

        st.image(
            image_bytes,
            caption=caption,
            use_container_width=True
        )

    except Exception:

        st.warning(
            "Gambar gagal ditampilkan."
        )

# =========================================================
# DASHBOARD
# =========================================================

if menu == "Dashboard":

    st.title("📊 Dashboard Buku Tamu")

    df = load_data()

    if not df.empty:

        df["tanggal"] = pd.to_datetime(
            df["tanggal"],
            errors="coerce"
        )

        today = datetime.now()

        total_tamu = len(df)

        harian = len(
            df[
                df["tanggal"].dt.date == today.date()
            ]
        )

        bulanan = len(
            df[
                (df["tanggal"].dt.month == today.month) &
                (df["tanggal"].dt.year == today.year)
            ]
        )

        tahunan = len(
            df[
                df["tanggal"].dt.year == today.year
            ]
        )

        c1, c2, c3, c4 = st.columns(4)

        c1.metric("👥 Total Tamu", total_tamu)
        c2.metric("📅 Hari Ini", harian)
        c3.metric("🗓️ Bulan Ini", bulanan)
        c4.metric("📆 Tahun Ini", tahunan)

        st.divider()

        st.subheader("📈 Grafik Kunjungan")

        grafik = (
            df.groupby(
                df["tanggal"].dt.date
            )
            .size()
            .reset_index(name="Jumlah")
        )

        grafik.columns = [
            "Tanggal",
            "Jumlah"
        ]

        st.line_chart(
            grafik.set_index("Tanggal")
        )

        st.divider()

        st.subheader("📋 Data Terbaru")

        preview_df = df.copy()

        preview_df["foto"] = "📸 Ada"
        preview_df["spt"] = "📄 Ada"

        st.dataframe(
            preview_df.tail(10),
            use_container_width=True,
            hide_index=True
        )

    else:

        st.info(
            "Belum ada data tamu."
        )

# =========================================================
# INPUT BUKU TAMU
# =========================================================

elif menu == "Input Buku Tamu":

    st.title("📝 Input Buku Tamu")

    with st.form(
        "form_tamu",
        clear_on_submit=True
    ):

        col1, col2 = st.columns(2)

        with col1:

            tanggal = datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )

            st.text_input(
                "Tanggal",
                value=tanggal,
                disabled=True
            )

            nama = st.text_input(
                "Nama Lengkap *"
            )

            opd = st.text_input(
                "Asal / OPD *"
            )

            nomor_hp = st.text_input(
                "Nomor HP *"
            )

        with col2:

            bidang = st.selectbox(
                "Bidang Tujuan *",
                [
                    "",
                    "Sekretariat",
                    "Bidang Litbang",
                    "Bidang Ekonomi",
                    "Bidang Sarana dan Prasarana Wilayah",
                    "Bidang Pemerintahan dan Sosial Budaya"
                ]
            )

            foto_tamu = st.file_uploader(
                "📷 Upload Foto Tamu",
                type=["jpg", "jpeg", "png"]
            )

            foto_spt = st.file_uploader(
                "📄 Upload Foto SPT",
                type=["jpg", "jpeg", "png"]
            )

        st.subheader("👁️ Preview Foto")

        p1, p2 = st.columns(2)

        with p1:

            if foto_tamu:
                st.image(
                    foto_tamu,
                    caption="Preview Foto Tamu",
                    use_container_width=True
                )

        with p2:

            if foto_spt:
                st.image(
                    foto_spt,
                    caption="Preview Foto SPT",
                    use_container_width=True
                )

        submit = st.form_submit_button(
            "💾 Simpan Data"
        )

        if submit:

            if nama.strip() == "":
                st.error("Nama wajib diisi.")

            elif opd.strip() == "":
                st.error("Asal / OPD wajib diisi.")

            elif nomor_hp.strip() == "":
                st.error("Nomor HP wajib diisi.")

            elif not nomor_hp.isdigit():
                st.error("Nomor HP harus berupa angka.")

            elif len(nomor_hp) < 10:
                st.error("Nomor HP tidak valid.")

            elif bidang == "":
                st.error("Bidang wajib dipilih.")

            elif foto_tamu is None:
                st.error("Foto tamu wajib diupload.")

            elif foto_spt is None:
                st.error("Foto SPT wajib diupload.")

            else:

                try:

                    foto_base64 = base64.b64encode(
                        foto_tamu.getvalue()
                    ).decode()

                    spt_base64 = base64.b64encode(
                        foto_spt.getvalue()
                    ).decode()

                    sheet.append_row([
                        tanggal,
                        nama,
                        opd,
                        nomor_hp,
                        bidang,
                        foto_base64,
                        spt_base64
                    ])

                    st.success(
                        "✅ Data berhasil disimpan."
                    )

                    st.balloons()

                    st.cache_data.clear()

                    st.rerun()

                except Exception as e:

                    st.error(
                        f"Gagal menyimpan data: {e}"
                    )

# =========================================================
# DAFTAR BUKU TAMU
# =========================================================

elif menu == "Daftar Buku Tamu":

    st.title("📑 Daftar Buku Tamu")

    df = load_data()

    if not df.empty:

        search = st.text_input(
            "🔍 Cari Nama / OPD"
        )

        if search:

            df = df[
                df.astype(str)
                .apply(
                    lambda x:
                    x.str.contains(
                        search,
                        case=False,
                        na=False
                    )
                )
                .any(axis=1)
            ]

        preview_df = df.copy()

        preview_df["foto"] = "📸 Ada"
        preview_df["spt"] = "📄 Ada"

        st.dataframe(
            preview_df,
            use_container_width=True,
            hide_index=True,
            height=400
        )

        st.divider()

        st.subheader("📷 Dokumentasi Tamu")

        for index, row in df.iterrows():

            with st.expander(
                f"👤 {row['nama']}"
            ):

                c1, c2 = st.columns(2)

                with c1:

                    st.write(
                        f"🏢 OPD : {row['opd']}"
                    )

                    st.write(
                        f"📞 HP : {row['nomor_hp']}"
                    )

                    st.write(
                        f"🏛️ Bidang : {row['bidang']}"
                    )

                    show_base64_image(
                        row["foto"],
                        "📸 Foto Tamu"
                    )

                with c2:

                    show_base64_image(
                        row["spt"],
                        "📄 Foto SPT"
                    )

        st.divider()

        st.subheader("📥 Download Data")

        col1, col2, col3, col4 = st.columns(4)

        # CSV

        csv = df.to_csv(
            index=False
        ).encode("utf-8")

        with col1:

            st.download_button(
                "⬇️ CSV",
                csv,
                "data_buku_tamu.csv",
                "text/csv"
            )

        # EXCEL

        excel_buffer = BytesIO()

        with pd.ExcelWriter(
            excel_buffer,
            engine="openpyxl"
        ) as writer:

            df.to_excel(
                writer,
                index=False
            )

        with col2:

            st.download_button(
                "⬇️ Excel",
                excel_buffer.getvalue(),
                "data_buku_tamu.xlsx"
            )

        # WORD

        with col3:

            if DOCX_AVAILABLE:

                try:

                    doc = Document()

                    doc.add_heading(
                        "Data Buku Tamu",
                        level=1
                    )

                    table = doc.add_table(
                        rows=1,
                        cols=len(df.columns)
                    )

                    hdr = table.rows[0].cells

                    for i, col in enumerate(df.columns):
                        hdr[i].text = col

                    for _, row in df.iterrows():

                        row_cells = table.add_row().cells

                        for i, value in enumerate(row):
                            row_cells[i].text = str(value)

                    word_buffer = BytesIO()

                    doc.save(word_buffer)

                    st.download_button(
                        "⬇️ Word",
                        word_buffer.getvalue(),
                        "data_buku_tamu.docx"
                    )

                except Exception as e:

                    st.warning(
                        f"Gagal membuat Word: {e}"
                    )

        # PDF

        with col4:

            if PDF_AVAILABLE:

                try:

                    pdf_buffer = BytesIO()

                    pdf_doc = SimpleDocTemplate(
                        pdf_buffer,
                        pagesize=letter
                    )

                    styles = getSampleStyleSheet()

                    elements = []

                    title = Paragraph(
                        "Data Buku Tamu",
                        styles["Heading1"]
                    )

                    elements.append(title)

                    elements.append(
                        Spacer(1, 12)
                    )

                    pdf_data = [
                        df.columns.tolist()
                    ] + df.values.tolist()

                    pdf_table = Table(
                        pdf_data,
                        repeatRows=1
                    )

                    pdf_table.setStyle(
                        TableStyle([
                            ('BACKGROUND', (0,0), (-1,0), colors.grey),
                            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
                            ('GRID', (0,0), (-1,-1), 1, colors.black),
                            ('FONTSIZE', (0,0), (-1,-1), 7),
                        ])
                    )

                    elements.append(
                        pdf_table
                    )

                    pdf_doc.build(
                        elements
                    )

                    st.download_button(
                        "⬇️ PDF",
                        pdf_buffer.getvalue(),
                        "data_buku_tamu.pdf"
                    )

                except Exception as e:

                    st.warning(
                        f"Gagal membuat PDF: {e}"
                    )

        st.divider()

        st.subheader("🗑️ Hapus Data")

        nama_hapus = st.selectbox(
            "Pilih Nama",
            df["nama"].tolist()
        )

        if st.button("🗑️ Hapus Data"):

            try:

                cell = sheet.find(
                    nama_hapus
                )

                if cell:

                    sheet.delete_rows(
                        cell.row
                    )

                    st.success(
                        "✅ Data berhasil dihapus."
                    )

                    st.cache_data.clear()

                    st.rerun()

            except Exception as e:

                st.error(
                    f"Gagal menghapus data: {e}"
                )

    else:

        st.info(
            "Belum ada data tamu."
        )

# =========================================================
# RINGKASAN STATISTIK
# =========================================================

elif menu == "Ringkasan Statistik":

    st.title("📊 Ringkasan Statistik")

    df = load_data()

    if not df.empty:

        df["tanggal"] = pd.to_datetime(
            df["tanggal"],
            errors="coerce"
        )

        statistik = (
            df.groupby(
                df["tanggal"].dt.date
            )
            .size()
            .reset_index(name="Jumlah")
        )

        statistik.columns = [
            "Tanggal",
            "Jumlah"
        ]

        st.subheader(
            "📈 Statistik Kunjungan"
        )

        st.bar_chart(
            statistik.set_index("Tanggal")
        )

        st.divider()

        c1, c2 = st.columns(2)

        c1.metric(
            "📸 Total Foto Tamu",
            len(df)
        )

        c2.metric(
            "📄 Total Foto SPT",
            len(df)
        )

        st.divider()

        st.subheader(
            "📷 Dokumentasi Visual Tamu"
        )

        for _, row in df.tail(10).iterrows():

            with st.expander(
                f"👤 {row['nama']}"
            ):

                col1, col2 = st.columns(2)

                with col1:

                    st.write(
                        f"🏢 OPD : {row['opd']}"
                    )

                    show_base64_image(
                        row["foto"],
                        "📸 Foto Tamu"
                    )

                with col2:

                    show_base64_image(
                        row["spt"],
                        "📄 Foto SPT"
                    )

    else:

        st.info(
            "Belum ada data statistik."
        )
```
